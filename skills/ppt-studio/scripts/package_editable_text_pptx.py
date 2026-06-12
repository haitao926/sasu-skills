#!/usr/bin/env python3
"""Package Roil visual plates with editable native PowerPoint text boxes."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import tempfile
import sys
from zipfile import ZipFile, ZIP_DEFLATED

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SLIDE_W = 13.333333
SLIDE_H = 7.5
DEFAULT_EAST_ASIA_FONT = "PingFang SC"
DEFAULT_LATIN_FONT = "Aptos"

ROLE_SIZES = {
    "title": 40,
    "subtitle": 22,
    "heading": 24,
    "body": 20,
    "label": 16,
    "page": 11,
}

ROLE_BOLD = {
    "title": True,
    "subtitle": True,
    "heading": True,
    "body": False,
    "label": False,
    "page": False,
}


def parse_color(raw: str | None, default: str = "#293241") -> RGBColor:
    value = (raw or default).strip().lstrip("#")
    if len(value) != 6:
        value = default.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def align_value(raw: str | None):
    if raw == "center":
        return PP_ALIGN.CENTER
    if raw == "right":
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def anchor_value(raw: str | None):
    if raw == "top":
        return MSO_ANCHOR.TOP
    if raw == "bottom":
        return MSO_ANCHOR.BOTTOM
    return MSO_ANCHOR.MIDDLE


def parse_rgba(raw: str | None, default: str = "#FFFFFF") -> tuple[RGBColor, float]:
    value = (raw or default).strip().lstrip("#")
    if len(value) == 6:
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)), 0.0
    if len(value) == 8:
        alpha = int(value[6:8], 16)
        transparency = max(0.0, min(1.0, 1 - (alpha / 255)))
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)), transparency
    color = default.strip().lstrip("#")
    return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)), 0.0


def ensure_typeface(rpr, tag: str, typeface: str) -> None:
    element = rpr.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        rpr.append(element)
    element.set("typeface", typeface)


def set_run_fonts(run, east_asia_font: str, latin_font: str) -> None:
    """Write every font slot WPS/PowerPoint may use for mixed Chinese text."""
    run.font.name = east_asia_font
    rpr = run._r.get_or_add_rPr()
    ensure_typeface(rpr, "a:latin", latin_font)
    ensure_typeface(rpr, "a:ea", east_asia_font)
    ensure_typeface(rpr, "a:cs", east_asia_font)


def patch_theme_fonts(pptx_path: Path, east_asia_font: str, latin_font: str) -> None:
    """Patch theme defaults because WPS may still consult them for CJK fallback."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        with ZipFile(pptx_path) as zin:
            zin.extractall(tmp)

        theme = tmp / "ppt" / "theme" / "theme1.xml"
        if theme.exists():
            try:
                from lxml import etree
            except ImportError:
                etree = None
            if etree is not None:
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                parser = etree.XMLParser(remove_blank_text=False)
                root = etree.parse(str(theme), parser)
                for font in root.xpath("//a:fontScheme/a:majorFont/a:latin | //a:fontScheme/a:minorFont/a:latin", namespaces=ns):
                    font.set("typeface", latin_font)
                for font in root.xpath("//a:fontScheme/a:majorFont/a:ea | //a:fontScheme/a:minorFont/a:ea", namespaces=ns):
                    font.set("typeface", east_asia_font)
                for font in root.xpath("//a:fontScheme/a:majorFont/a:cs | //a:fontScheme/a:minorFont/a:cs", namespaces=ns):
                    font.set("typeface", east_asia_font)
                theme.write_bytes(etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=False))

        with ZipFile(pptx_path, "w", ZIP_DEFLATED) as zout:
            for file in tmp.rglob("*"):
                if file.is_file():
                    zout.write(file, file.relative_to(tmp).as_posix())


def add_container(slide, block: dict, x, y, w, h) -> None:
    kind = str(block.get("container_shape", "rounded"))
    if kind == "rect":
        shape_type = MSO_AUTO_SHAPE_TYPE.RECTANGLE
    else:
        shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    fill_color, fill_transparency = parse_rgba(block.get("container_fill"), "#FFFFFF")
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = fill_transparency
    line_color, line_transparency = parse_rgba(block.get("container_line"), "#A9C7E8")
    shape.line.color.rgb = line_color
    shape.line.transparency = line_transparency
    shape.line.width = Pt(float(block.get("container_line_width", 1.5)))
    if block.get("container_shadow") is False:
        shape.shadow.inherit = False


def add_textbox(slide, block: dict) -> None:
    x = Inches(float(block["x"]) * SLIDE_W)
    y = Inches(float(block["y"]) * SLIDE_H)
    w = Inches(float(block["w"]) * SLIDE_W)
    h = Inches(float(block["h"]) * SLIDE_H)
    role = str(block.get("role", "body"))
    if block.get("container"):
        add_container(slide, block, x, y, w, h)
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(float(block.get("margin_left", 0.02)))
    frame.margin_right = Inches(float(block.get("margin_right", 0.02)))
    frame.margin_top = Inches(float(block.get("margin_top", 0.01)))
    frame.margin_bottom = Inches(float(block.get("margin_bottom", 0.01)))
    frame.vertical_anchor = anchor_value(block.get("vertical_anchor"))
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align_value(block.get("align"))
    paragraph.text = str(block.get("text", ""))
    if "line_spacing" in block:
        paragraph.line_spacing = float(block["line_spacing"])
    if "space_before" in block:
        paragraph.space_before = Pt(float(block["space_before"]))
    if "space_after" in block:
        paragraph.space_after = Pt(float(block["space_after"]))
    east_asia_font = str(block.get("east_asia_font") or block.get("font") or DEFAULT_EAST_ASIA_FONT)
    latin_font = str(block.get("latin_font") or DEFAULT_LATIN_FONT)
    for run in paragraph.runs:
        set_run_fonts(run, east_asia_font, latin_font)
        run.font.size = Pt(float(block.get("font_size", ROLE_SIZES.get(role, 18))))
        run.font.bold = bool(block.get("bold", ROLE_BOLD.get(role, False)))
        run.font.color.rgb = parse_color(block.get("color"))
        run.font._element.set("dirty", "0")
    box.fill.background()
    box.line.fill.background()


def package(run_dir: Path, spec: dict, out_path: Path) -> None:
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("spec must contain a non-empty slides array")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    for index, slide_spec in enumerate(slides, start=1):
        if not isinstance(slide_spec, dict):
            raise ValueError(f"slide {index} is not an object")
        plate = run_dir / str(slide_spec["visual_plate"])
        if not plate.exists():
            raise FileNotFoundError(f"visual plate not found: {slide_spec['visual_plate']}")
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(plate), 0, 0, width=prs.slide_width, height=prs.slide_height)
        for block in slide_spec.get("text_blocks", []):
            add_textbox(slide, block)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    patch_theme_fonts(
        out_path,
        str(spec.get("east_asia_font", DEFAULT_EAST_ASIA_FONT)),
        str(spec.get("latin_font", DEFAULT_LATIN_FONT)),
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Package editable PPT text over Roil visual plates.")
    parser.add_argument("--run-dir", required=True)
    parser.add_argument("--spec", required=True)
    parser.add_argument("--out", default="exports/presentation-editable.pptx")
    args = parser.parse_args(argv)

    run_dir = Path(args.run_dir)
    spec_path = Path(args.spec)
    if not spec_path.is_absolute():
        spec_path = run_dir / spec_path
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out_path = Path(args.out)
    if not out_path.is_absolute():
        out_path = run_dir / out_path
    package(run_dir, spec, out_path)
    print(json.dumps({"ok": True, "pptx": str(out_path), "slides": len(spec["slides"])}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
